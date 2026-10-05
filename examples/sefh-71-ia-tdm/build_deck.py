"""Maqueta la ponencia sobre la plantilla SEFH «71-congreso-plantilla-ponencias.pptx».

    python3 build_deck.py <plantilla.pptx> <salida.pptx>

Trabaja sobre el propio paquete de la plantilla (patrón, diseño, tema y fondos),
de modo que el resultado se edita en PowerPoint como cualquier archivo SEFH.
Coordenadas en píxeles de un lienzo de 1920 x 1080 (rejilla de 12 columnas,
márgenes de 80 px y medianil de 24 px).
"""

import copy
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
ASSETS = HERE / "assets"

# --- Lienzo y rejilla --------------------------------------------------------
SLIDE_W, SLIDE_H = 12192000, 6858000          # 13,333 x 7,5 in (16:9)
EMU = SLIDE_W / 1920                          # EMU por píxel del lienzo
MARGIN, GUTTER = 80, 24
COLW = (1920 - 2 * MARGIN - 11 * GUTTER) / 12


def px(v):
    return Emu(int(round(v * EMU)))


def col(n):
    """x del borde izquierdo de la columna n (1-12)."""
    return MARGIN + (n - 1) * (COLW + GUTTER)


def cols(a, b):
    """(x, ancho) del tramo de columnas a..b."""
    return col(a), col(b) + COLW - col(a)


# Franjas verticales
TITLE_Y, TITLE_H = 118, 150        # bajo la banda de la plantilla
CONTENT_Y, CONTENT_B = 290, 985
SOURCE_Y, SOURCE_H = 998, 38       # pie de fuente (dentro del marco, sobre los triángulos)

# --- Colores (de la plantilla) -----------------------------------------------
INK = "263445"
GRAY = "5E6B78"
LIGHT = "F3F6F9"
RULE = "C9D2DB"
WHITE = "FFFFFF"
LIMA, LIMA_D, LIMA_T = "C3BF24", "77730E", "F4F3D6"
VERDE, VERDE_D, VERDE_T = "5BB46C", "2F8541", "E7F4E9"
TURQ, TURQ_D, TURQ_T = "00B3B2", "007F7E", "E0F5F5"
MORADO, MORADO_D, MORADO_T = "AD3B94", "8E2F79", "F8EAF4"
AZUL, AZUL_D, AZUL_T = "4A73A0", "2F5480", "E4EBF4"
ROJO, ROJO_T = "D42A3B", "FCE4E6"
AMBAR = "F0A500"
XRAY_BG, XRAY = "0E1B2C", "A9DBFF"

NB = " "   # espacio de no separación: cifra y unidad nunca se parten


def text_width(t, pt=20, bold=False):
    """Ancho aproximado del texto en px del lienzo (métrica de Calibri vía Carlito)."""
    try:
        from PIL import ImageFont
        f = ("/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf" if bold else
             "/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf")
        return ImageFont.truetype(f, pt * 20).getlength(t) / 10
    except OSError:
        return len(t) * pt * (1.0 if bold else 0.95)


def rgb(h):
    return RGBColor.from_string(h)


# --- Primitivas ---------------------------------------------------------------

def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _name(shape, name):
    if name:
        shape._element.xpath("./*[1]/p:cNvPr")[0].set("name", name)


def fill_text(tf, paras, size=20, color=INK, bold=False, italic=False, align="l",
              anchor="t", spacing=None, margins=(0, 0, 0, 0)):
    """Rellena un marco de texto.

    paras: lista de párrafos. Cada párrafo es un str, una lista de runs
    (str | (texto, {opciones})) o un dict {"runs": [...], "align", "after", "size", ...}.
    """
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    l, t, r, b = margins
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (px(l), px(t), px(r), px(b))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        popts = {}
        if isinstance(para, dict):
            popts = para
            runs = para["runs"]
        elif isinstance(para, str):
            runs = [para]
        else:
            runs = para
        if isinstance(runs, str):
            runs = [runs]
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[
            popts.get("align", align)]
        if "after" in popts:
            p.space_after = Pt(popts["after"])
        if "before" in popts:
            p.space_before = Pt(popts["before"])
        ls = popts.get("spacing", spacing)
        if ls:
            p.line_spacing = ls
        for run in runs:
            text, o = (run, {}) if isinstance(run, str) else run
            o = {**{k: popts[k] for k in ("size", "color", "bold", "italic") if k in popts}, **o}
            r_ = p.add_run()
            r_.text = text
            f = r_.font
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
            rpr = r_._r.get_or_add_rPr()
            rpr.set("lang", "es-ES")
            if o.get("sub"):
                rpr.set("baseline", "-25000")
    return tf


def text(slide, x, y, w, h, paras, name=None, **kw):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    _name(tb, name)
    fill_text(tb.text_frame, paras, **kw)
    return tb


def box(slide, x, y, w, h, fill=None, line=None, lw=1.5, dash=None, radius=None,
        kind=MSO_SHAPE.RECTANGLE, name=None, paras=None, **tkw):
    if radius is not None and kind == MSO_SHAPE.RECTANGLE:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE
    sp = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    _strip_style(sp)
    _name(sp, name)
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
        if dash:
            sp.line.dash_style = dash
    else:
        sp.line.fill.background()
    if radius is not None:
        sp.adjustments[0] = radius / min(w, h)
    if paras is not None:
        fill_text(sp.text_frame, paras, **tkw)
    return sp


def _line_ends(ln, head=None, tail=None):
    for tag, kind in (("a:headEnd", head), ("a:tailEnd", tail)):
        if kind:
            el = etree.SubElement(ln, qn(tag))
            el.set("type", kind)
            el.set("w", "med")
            el.set("len", "med")


def line(slide, x1, y1, x2, y2, color=INK, lw=2, dash=None, head=None, tail=None, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    _strip_style(c)
    _name(c, name)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = dash
    _line_ends(c.line._get_or_add_ln(), head, tail)
    return c


def curve(slide, p0, c1, c2, p1, color=INK, lw=2.5, tail="triangle", dash=None, name="Flecha curva"):
    """Flecha curva (Bézier cúbica) como forma libre editable."""
    xs = [p[0] for p in (p0, c1, c2, p1)]
    ys = [p[1] for p in (p0, c1, c2, p1)]
    x0, y0 = min(xs), min(ys)
    w, h = max(max(xs) - x0, 1), max(max(ys) - y0, 1)

    def pt(p):
        return f'<a:pt x="{int((p[0] - x0) * EMU)}" y="{int((p[1] - y0) * EMU)}"/>'

    sid = slide.shapes._next_shape_id
    dash_xml = f'<a:prstDash val="{dash}"/>' if dash else ""
    xml = f"""
<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{int(x0 * EMU)}" y="{int(y0 * EMU)}"/><a:ext cx="{int(w * EMU)}" cy="{int(h * EMU)}"/></a:xfrm>
    <a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>
      <a:pathLst><a:path w="{int(w * EMU)}" h="{int(h * EMU)}" fill="none">
        <a:moveTo>{pt(p0)}</a:moveTo><a:cubicBezTo>{pt(c1)}{pt(c2)}{pt(p1)}</a:cubicBezTo>
      </a:path></a:pathLst></a:custGeom>
    <a:noFill/>
    <a:ln w="{int(Pt(lw))}" cap="rnd"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>{dash_xml}
      <a:round/><a:tailEnd type="{tail}" w="med" len="med"/></a:ln>
  </p:spPr>
</p:sp>"""
    el = etree.fromstring(xml)
    slide.shapes._spTree.append(el)
    return el


def image(slide, name, x, y, w, h, align="c", valign="m", rotation=0, label=None):
    """Coloca un PNG ajustado dentro de la caja sin deformarlo."""
    path = ASSETS / f"{name}.png"
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    dw, dh = iw * s, ih * s
    dx = {"l": 0, "c": (w - dw) / 2, "r": w - dw}[align]
    dy = {"t": 0, "m": (h - dh) / 2, "b": h - dh}[valign]
    pic = slide.shapes.add_picture(str(path), px(x + dx), px(y + dy), px(dw), px(dh))
    _name(pic, label or name)
    pic._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", label or name)
    if rotation:
        pic.rotation = rotation
    return pic


def title(slide, paras, size=32, align="l", anchor="m", y=TITLE_Y, h=TITLE_H):
    return text(slide, MARGIN, y, 1920 - 2 * MARGIN, h, paras, size=size, color=INK,
                align=align, anchor=anchor, spacing=0.92, name="Titular")


def source(slide, s, h=SOURCE_H, y=SOURCE_Y):
    return text(slide, MARGIN, y, 1920 - 2 * MARGIN, h, s, size=12, color=GRAY,
                anchor="b", name="Fuente")


def icon_disc(slide, icon, cx, cy, d, fill, ring=None):
    """Icono de línea sobre un círculo de color suave."""
    box(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, line=ring, lw=1.5,
        kind=MSO_SHAPE.OVAL, name=f"Fondo {icon}")
    image(slide, icon, cx - d * 0.34, cy - d * 0.34, d * 0.68, d * 0.68)


# --- Diapositivas -------------------------------------------------------------

def slide_cover(s):
    """Rellena los cuadros de texto de la portada de la plantilla."""
    sp = {shape.name: shape for shape in s.shapes}

    def retext(shape, paras, x, y, w, h):
        shape.left, shape.top, shape.width, shape.height = px(x), px(y), px(w), px(h)
        tf = shape.text_frame
        bp = tf._txBody.find(qn("a:bodyPr"))
        for child in list(bp):
            bp.remove(child)                     # sin autoajuste: tamaño fijo
        p0 = tf.paragraphs[0]
        rpr_tpl = copy.deepcopy(p0.runs[0]._r.find(qn("a:rPr")))
        for p in list(tf._txBody.findall(qn("a:p")))[1:]:
            tf._txBody.remove(p)
        for r in list(p0._p.findall(qn("a:r"))):
            p0._p.remove(r)
        for i, (txt, opts) in enumerate(paras):
            p = p0 if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = txt
            rpr = copy.deepcopy(rpr_tpl)
            old = r._r.find(qn("a:rPr"))
            if old is not None:
                r._r.remove(old)
            r._r.insert(0, rpr)
            rpr.set("lang", "es-ES")
            rpr.set("sz", str(int(opts["size"] * 100)))
            if "bold" in opts:
                rpr.set("b", "1" if opts["bold"] else "0")
            if "italic" in opts:
                rpr.set("i", "1" if opts["italic"] else "0")
            if "color" in opts:
                for f in rpr.findall(qn("a:solidFill")):
                    rpr.remove(f)
                sf = etree.Element(qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr")).set("val", opts["color"])
                # solidFill va después de a:ln y antes de a:effectLst
                ln = rpr.find(qn("a:ln"))
                if ln is not None:
                    ln.addnext(sf)
                else:
                    rpr.insert(0, sf)
            if "spacing" in opts:
                p.line_spacing = opts["spacing"]
        tf.word_wrap = True

    retext(sp["TextBox 12"], [
        ("CURSO PRECONGRESO ·", {"size": 24, "bold": True}),
        ("INTELIGENCIA ARTIFICIAL APLICADA A LA FARMACOCINÉTICA CLÍNICA", {"size": 24, "bold": True}),
    ], 100, 352, 1720, 110)
    retext(sp["TextBox 4"], [
        ("«Qué puede aportar realmente la IA a la TDM y a la farmacocinética clínica»",
         {"size": 36, "bold": True, "spacing": 0.95}),
    ], 200, 486, 1520, 180)
    retext(sp["TextBox 6"], [("ALFREDO MONTERO DELGADO", {"size": 24, "bold": True, "color": INK})],
           260, 712, 1400, 50)
    retext(sp["TextBox 7"], [
        ("Farmacéutico especialista en Farmacia Hospitalaria · Hospital Universitario "
         "Nuestra Señora de Candelaria · Grupo DIGIFHAR (SEFH)", {"size": 20, "color": GRAY}),
    ], 260, 772, 1400, 110)


def slide_hook(s):
    title(s, [[("ChatGPT sabe hablar de farmacocinética, pero no ", {}),
               ("hacer", {"bold": True, "color": TURQ_D}),
               (" farmacocinética.", {})]])

    # Ilustración: columnas 1-7
    x0, w0 = cols(1, 7)
    floor_y = 972
    line(s, x0 + 10, floor_y, x0 + w0, floor_y, color=RULE, lw=2, name="Suelo")
    # Pizarra con fórmulas impecables (texto editable)
    bx, by, bw, bh = 480, 300, 610, 318
    box(s, bx, by, bw, bh, fill="31454F", line=INK, lw=3, radius=10, name="Pizarra")
    box(s, bx + 40, by + bh, bw - 80, 12, fill="B9C3CD", line=INK, lw=1.5, name="Repisa")
    box(s, bx + 430, by + bh - 10, 38, 10, fill=WHITE, line=INK, lw=1, name="Tiza")
    chalk = "F4F1E6"
    text(s, bx + 50, by + 24, bw - 90, bh - 40, [
        ["CL = k", ("e", {"sub": True}), " · V", ("d", {"sub": True})],
        ["t½ = 0,693 / k", ("e", {"sub": True})],
        ["V", ("d", {"sub": True}), " = Dosis / C", ("0", {"sub": True})],
        ["k", ("e", {"sub": True}), " = ln(C", ("1", {"sub": True}), " / C", ("2", {"sub": True}),
         ") / Δt"],
    ], size=26, color=chalk, italic=True, spacing=1.0, anchor="m", name="Fórmulas")
    # Loro muy seguro de sí mismo
    image(s, "loro-seguro", 90, 400, 380, floor_y - 400 + 6, align="c", valign="b",
          label="Loro robótico con bata señalando la pizarra")
    # Bocadillo
    bub = box(s, 520, 670, 560, 140, fill=WHITE, line=INK, lw=2.25,
              kind=MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, name="Bocadillo",
              paras=["aclaramiento… volumen de distribución… semivida…"],
              size=22, italic=True, color=INK, anchor="m", margins=(28, 10, 28, 10))
    bub.adjustments[0] = (372 - (520 + 280)) / 560
    bub.adjustments[1] = (566 - (670 + 70)) / 140
    # Calculadora boca abajo en el suelo
    image(s, "icono-calculadora", 780, floor_y - 112, 150, 112, valign="b", rotation=180,
          label="Calculadora boca abajo")

    # Tarjeta: columnas 8-12, centrada en vertical
    cx, cw = cols(8, 12)
    ch = 640
    cy = CONTENT_Y + (CONTENT_B - CONTENT_Y - ch) / 2
    box(s, cx, cy, cw, ch, fill=MORADO_T, radius=18, name="Tarjeta dato")
    text(s, cx + 48, cy + 24, cw - 96, 214, f"37{NB}%", size=110, bold=True, color=MORADO,
         anchor="m", name="Cifra")
    text(s, cx + 48, cy + 250, cw - 96, 214,
         f"de las predicciones de valle de vancomicina de ChatGPT quedaron a ±2{NB}mg/L "
         f"del valor medido", size=22, color=INK, name="Lectura")
    line(s, cx + 48, cy + 482, cx + cw - 48, cy + 482, color="DDB7D3", lw=1.5)
    text(s, cx + 48, cy + 500, cw - 96, 110,
         f"717 predicciones · 239 ingresos en UCI · Grok: 29{NB}%", size=20, color=GRAY,
         name="Detalle")

    source(s, "Ishaqui et al., Expert Rev Anti Infect Ther, 2026")


def slide_not_ai(s):
    title(s, "No todo lo que calcula es inteligencia artificial.")
    box(s, MARGIN, 288, 1760, 70, fill=LIGHT, radius=35, name="Pregunta guía",
        paras=["¿Quién decide la forma del modelo: una persona o los datos?"],
        size=24, bold=True, color=TURQ_D, align="c", anchor="m")

    lx, lw_ = cols(1, 6)
    rx, rw = cols(7, 12)
    head_y = 384
    text(s, lx, head_y, lw_, 46, "La decide una persona", size=24, bold=True, color=AZUL_D)
    text(s, rx, head_y, rw, 46, "La aprenden los datos", size=24, bold=True, color=MORADO_D)
    line(s, 960, 392, 960, 965, color=RULE, lw=1.5, dash=MSO_LINE.DASH, name="Separador")

    cards = [
        ("icono-curva-color", "Modelos poblacionales y estimación bayesiana"),
        ("icono-calculadora", "Calculadoras y nomogramas"),
        ("icono-engranaje", "Automatizaciones por reglas"),
    ]
    cy, chh, gap = 450, 150, 24
    for i, (ic, label) in enumerate(cards):
        y = cy + i * (chh + gap)
        box(s, lx, y, lw_ - 24, chh, fill=LIGHT, radius=16, name=f"Tarjeta {i + 1}")
        icon_disc(s, ic, lx + 90, y + chh / 2, 112, WHITE)
        text(s, lx + 176, y, lw_ - 24 - 200, chh, label, size=22, color=INK, anchor="m",
             bold=True)

    # Círculos concéntricos: de lima (fuera) a morado (dentro)
    ccx, ccy = rx + 222, 712
    rings = [(222, LIMA, "Aprendizaje automático", LIMA_D),
             (150, "3A8AA6", ["Aprendizaje profundo", "(redes neuronales)"], "2B6C85"),
             (80, MORADO, ["Modelos de lenguaje", "(IA generativa)"], MORADO_D)]
    for r, c, _, _ in rings:
        box(s, ccx - r, ccy - r, 2 * r, 2 * r, fill=c, line=WHITE, lw=3, kind=MSO_SHAPE.OVAL,
            name="Círculo")
    lab_x = rx + 452
    lab_w = rx + rw - lab_x
    # (punto dentro de cada anillo, y de la etiqueta)
    anchors = [((ccx + 120, ccy - 140), ccy - 140),
               ((ccx + 115, ccy - 6), ccy - 6),
               ((ccx + 30, ccy + 40), ccy + 136)]
    for (r, c, label, dark), ((dx, dy), ly) in zip(rings, anchors):
        box(s, dx - 9, dy - 9, 18, 18, fill=WHITE, line=INK, lw=2, kind=MSO_SHAPE.OVAL,
            name="Marca")
        if abs(ly - dy) < 1:
            line(s, dx + 9, dy, lab_x - 12, ly, color=INK, lw=1.5)
        else:
            line(s, dx + 6, dy + 6, lab_x - 70, ly, color=INK, lw=1.5)
            line(s, lab_x - 70, ly, lab_x - 12, ly, color=INK, lw=1.5)
        paras = [label] if isinstance(label, str) else label     # salto de línea, mismo texto
        text(s, lab_x, ly - 52, lab_w, 104, paras, size=20, bold=True, color=dark, anchor="m")

    source(s, "Keutzer et al., Pharmaceutics, 2022")


def slide_thirty_years(s):
    title(s, "Treinta años de estudios: predice parecido al bayesiano, pero apenas ha llegado "
             "al día a día.")
    # Línea de tiempo: flecha fina con dos hitos
    ty = 312
    line(s, MARGIN, ty, 1840, ty, color=INK, lw=2, tail="triangle", name="Línea de tiempo")
    lx, lw_ = cols(1, 6)
    rx, rw = cols(7, 12)
    for (x, w, al, yr, rest) in (
        (lx, lw_, "l", "1995", " · Red neuronal frente a NONMEM (gentamicina, 111 pacientes)"),
        (rx, rw, "r", "2025", " · 58 estudios: el aprendizaje automático rinde igual o mejor "
                             "que los modelos poblacionales"),
    ):
        dot_x = x + 14 if al == "l" else x + w - 60
        box(s, dot_x - 13, ty - 13, 26, 26, fill=WHITE, line=INK, lw=3, kind=MSO_SHAPE.OVAL,
            name=f"Hito {yr}")
        text(s, x, ty + 26, w, 130, [[(yr, {"bold": True, "size": 24}), (rest, {})]],
             size=20, color=INK, align=al, name=f"Texto {yr}")

    # Tarjetas
    cy, chh = 476, 366
    box(s, lx, cy, lw_, chh, fill=LIGHT, radius=18, name="Tarjeta sin concentraciones")
    image(s, "icono-gotero", lx + 32, cy + 28, 60, 84, label="Gotero vacío")
    text(s, lx + 112, cy + 30, lw_ - 140, 80, "Sin concentraciones del paciente", size=24,
         bold=True, color=INK, anchor="m")
    text(s, lx + 40, cy + 136, lw_ - 80, 200, [[
        ("El modelo híbrido mejora un ", {}), (f"17{NB}%", {"bold": True, "color": MORADO}),
        (" el error del bayesiano · vancomicina en sepsis, 4059 pacientes", {})]],
        size=22, color=INK)

    box(s, rx, cy, rw, chh, fill=LIGHT, radius=18, name="Tarjeta con concentraciones")
    image(s, "icono-tubo", rx + 32, cy + 28, 60, 84, label="Tubo de análisis")
    text(s, rx + 112, cy + 30, rw - 140, 80, "Con concentraciones del paciente", size=24,
         bold=True, color=INK, anchor="m")
    for i, (c, dark, runs) in enumerate((
        (VERDE, VERDE_D, [("Gana el bayesiano: ", {"bold": True}),
                          (f"13{NB}% de error frente{NB}a{NB}34{NB}%", {})]),
        (MORADO, MORADO_D, [("Gana el aprendizaje automático: ", {"bold": True}),
                            ("ABC de tacrolimus en 6 series externas", {})]),
    )):
        yy = cy + 132 + i * 112
        box(s, rx + 40, yy + 10, 20, 20, fill=c, kind=MSO_SHAPE.OVAL, name="Marca color")
        text(s, rx + 76, yy, rw - 116, 100, [runs], size=22, color=dark)

    box(s, MARGIN, 862, 1760, 70, fill=WHITE, line=GRAY, lw=1.5, dash=MSO_LINE.DASH, radius=35,
        name="Franja", paras=["Casi todo retrospectivo · La IA y los híbridos siguen en investigación"],
        size=22, bold=True, color=GRAY, align="c", anchor="m")

    source(s, "Brier et al., Pharm Res, 1995 · Methaneethorn et al., Clin Pharmacokinet, 2025 · "
              "Chen et al., Microbiol Spectr, 2025 · Woillard et al., Clin Pharmacol Ther, 2021 · "
              "Altynova et al., Pharmaceuticals, 2026", h=56, y=980)


def slide_ladder(s):
    title(s, "Predecir bien no es lo mismo que mejorar al paciente.")
    # Escalera: cinco bandas apiladas que suben en diagonal (columnas 1-8)
    x_end = col(8) + COLW
    base = 965
    rise, run = 128, 150
    x0 = MARGIN
    steps = [
        ("Estudio retrospectivo", None),
        ("Validación externa", "10 de 115 estudios de tacrolimus"),
        ("Estudio prospectivo", "pocos, pequeños, sin comparador (CURATE.AI: 10 pacientes)"),
        ("Ensayo con desenlace de proceso", None),
        ("Ensayo con desenlace clínico", None),
    ]
    tints = ["E9EEF3", "DEE5EC", "D3DCE5", "C8D3DE", "BDCAD7"]
    for i, (name, detail) in enumerate(steps):
        top = base - (i + 1) * rise
        x = x0 + i * run
        box(s, x, top, x_end - x, rise, fill=tints[i], line=WHITE, lw=2,
            name=f"Peldaño {i + 1}")
        runs = [(name, {"bold": True})]
        if detail:
            runs += [(": " + detail, {})]
        text(s, x + 22, top + 6, x_end - x - 44, rise - 12, [runs], size=20, color=INK,
             anchor="m", name=f"Rótulo peldaño {i + 1}")

    # Marcador morado: mancha densa en 1, se estrecha en 2, un punto en 3
    def tread(i):
        return base - (i + 1) * rise

    import random
    rnd = random.Random(11)
    for i, n, spread in ((0, 78, 54), (1, 14, 30), (2, 1, 0)):
        cxp = x0 + i * run + run / 2
        if n > 1:
            box(s, cxp - spread - 16, tread(i) - spread * 1.35 - 26, 2 * spread + 32,
                spread * 1.35 + 30, fill=MORADO_T, kind=MSO_SHAPE.OVAL, name="Mancha")
        for _ in range(n):
            dx = rnd.uniform(-spread, spread)
            dy = -abs(rnd.gauss(0, spread * 0.5)) - 10
            r = rnd.uniform(5, 8.5) if n > 1 else 11
            box(s, cxp + dx - r, tread(i) + dy - r, 2 * r, 2 * r, fill=MORADO,
                line=WHITE, lw=0.75, kind=MSO_SHAPE.OVAL, name="Estudio")
    text(s, x0, 460, 290, 100, "Aprendizaje automático", size=22, bold=True, color=MORADO,
         anchor="b", name="Rótulo marcador")
    line(s, x0 + 50, 566, x0 + 50, 730, color=MORADO, lw=2, dash=MSO_LINE.ROUND_DOT,
         name="Guía marcador")

    # Tarjeta verde unida al quinto peldaño
    cx, cw = cols(9, 12)
    cy, chh = 296, 650
    head_h = 196
    box(s, cx, cy, cw, chh, fill=VERDE_T, line=VERDE, lw=2.5, radius=18, name="Tarjeta TDM")
    box(s, cx, cy, cw, head_h, fill=VERDE_D, radius=18, name="Cabecera tarjeta")
    box(s, cx, cy + head_h - 40, cw, 40, fill=VERDE_D, name="Cabecera tarjeta (base)")
    text(s, cx + 32, cy + 16, cw - 64, head_h - 32,
         "Dosificación individualizada de antimicrobianos (TDM o MIPD)", size=22, bold=True,
         color=WHITE, anchor="m")
    text(s, cx + 32, cy + head_h + 16, cw - 64, 44, "10 ensayos · 1241 pacientes", size=20,
         color=VERDE_D, bold=True)
    rows_ = [
        [("Fracaso terapéutico ", {}), (f"−30{NB}%", {"bold": True}), (" (RR 0,70)", {})],
        [("Nefrotoxicidad ", {}), (f"−45{NB}%", {"bold": True}), (" (RR 0,55)", {})],
        [("Mortalidad: sin cambio significativo", {})],
    ]
    for i, r in enumerate(rows_):
        yy = cy + head_h + 76 + i * 122
        if i:
            line(s, cx + 32, yy - 10, cx + cw - 32, yy - 10, color="B9DDBF", lw=1.25)
        text(s, cx + 32, yy, cw - 64, 108, [r], size=22, color=INK, anchor="m")
    # Unión: del quinto peldaño a la tarjeta
    y5 = tread(4) + rise / 2
    line(s, x_end - 6, y5, cx, y5, color=VERDE, lw=4, name="Unión peldaño 5")
    box(s, x_end - 18, y5 - 12, 24, 24, fill=VERDE, line=WHITE, lw=2, kind=MSO_SHAPE.OVAL,
        name="Nodo peldaño 5")

    source(s, "Amooei et al., Pharmaceutics, 2026 · Blasiak et al., NPJ Precis Oncol, 2025 · "
              "Sanz-Codina et al., Clin Microbiol Infect, 2023")


# Geometría común del esquema de las cinco piezas (diapositivas 6 y 7)
FRAME_X, FRAME_W = cols(2, 11)
PCOLS = [(FRAME_X + 20, 400), (FRAME_X + 440, 582), (FRAME_X + FRAME_W - 420, 400)]
PHARM = (902, 290, 116, 112)          # x, y, ancho, alto del farmacéutico


def _center(i):
    x, w = PCOLS[i]
    return x + w / 2


def scheme_art(s, variant, img_y, img_h, arrow_color):
    """Dibujos y flechas del esquema (ficha, loro, herramientas)."""
    image(s, f"ficha-caso-{variant}", _center(0) - 70, img_y, 140, img_h, label="Ficha del caso")
    image(s, f"loro-{variant}", _center(1) - 60, img_y - 8, 120, img_h + 14,
          label="Loro robótico (modelo de lenguaje)")
    step = img_h / 3
    for i, ic in enumerate(("curva", "campana", "red")):
        image(s, f"icono-{ic}-{variant}", _center(2) - 44, img_y + i * step + 2, 88, step - 6,
              label=f"Icono {ic}")
    ay = img_y + img_h / 2
    line(s, _center(0) + 90, ay, _center(1) - 70, ay, color=arrow_color, lw=2.5,
         tail="triangle", name="Contexto → modelo")
    line(s, _center(1) + 70, ay, _center(2) - 70, ay, color=arrow_color, lw=2.5,
         head="triangle", tail="triangle", name="Modelo ↔ herramientas")


def slide_pieces(s):
    title(s, "El fallo no es de la IA: es de usar una sola pieza.")
    fy, fh = 420, 456
    # Farmacéutico, fuera del marco y por encima
    px_, py_, pw_, ph_ = PHARM
    image(s, "farmaceutico-color", px_, py_, pw_, ph_, label="Farmacéutico con pluma")
    text(s, px_ + pw_ + 18, py_, 720, ph_, [
        [("Farmacéutico", {"bold": True, "color": AZUL_D}), (" · supervisa y aporta criterio", {})],
        [("firma del informe", {"color": GRAY})],
    ], size=20, color=INK, anchor="m")
    line(s, 960, py_ + ph_ + 2, 960, fy, color=AZUL, lw=2, dash=MSO_LINE.ROUND_DOT,
         tail="triangle", name="Supervisión")

    # Arnés: marco discontinuo con su rótulo en el borde
    box(s, FRAME_X, fy, FRAME_W, fh, line=INK, lw=2.5, dash=MSO_LINE.DASH, radius=24, name="Arnés")
    box(s, FRAME_X + 26, fy - 24, 700, 96, fill=WHITE, name="Rótulo arnés",
        paras=[[("Arnés (el protocolo)", {"bold": True})],
               [("procedimiento normalizado de trabajo", {"color": GRAY})]],
        size=20, color=INK, anchor="t", margins=(14, 2, 14, 0))
    box(s, 960 - 220, fy + fh - 22, 440, 44, fill=WHITE, name="Borde arnés",
        paras=["reglas · límites · registro"], size=20, italic=True, color=INK, align="c",
        anchor="m")

    img_y, img_h = fy + 92, 150
    scheme_art(s, "color", img_y, img_h, INK)
    lab_y = img_y + img_h + 16
    labels = [
        [[("Contexto", {"bold": True, "color": TURQ_D}), (" · lo que le das del caso", {})],
         [("historia clínica", {"color": GRAY})]],
        [[("Modelo de lenguaje", {"bold": True}), (" · predice texto", {})],
         [("lee, llama a la herramienta y redacta el borrador", {"color": GRAY})]],
        [[("Herramientas", {"bold": True, "color": VERDE_D}),
          (" · código y programas validados", {})],
         [("programa bayesiano", {"color": GRAY})]],
    ]
    for (x, w), lab in zip(PCOLS, labels):
        text(s, x, lab_y, w, 150, lab, size=20, color=INK, align="c")

    text(s, MARGIN, 912, 1760, 60, [[
        ("El modelo de lenguaje no da siempre la misma respuesta; ", {}),
        ("la herramienta validada, sí.", {"bold": True, "color": VERDE_D})]],
        size=24, color=INK, align="c", anchor="m", name="Frase de pie")
    source(s, "Pritchard-Bell et al., CPT Pharmacometrics Syst Pharmacol, 2026")


def slide_xray(s):
    title(s, "La lección de TDM-AID: cada tarea, a la pieza que mejor la hace.")
    box(s, MARGIN, 282, 1760, 700, fill=XRAY_BG, radius=20, name="Fondo radiografía")
    fy, fh = 420, 420
    soft = "DDE7F0"

    def tag(x, y, w, label, color, fg=XRAY_BG, h=46, name="Etiqueta"):
        if w is None:
            w = text_width(label, 20, True) + 60
        return box(s, x, y, w, h, fill=color, radius=22, name=name, paras=[label], size=20,
                   bold=True, color=fg, align="c", anchor="m", margins=(14, 0, 14, 0))

    def ctag(i, y, w, label, color, fg=XRAY_BG, h=46, name="Etiqueta"):
        w = w or text_width(label, 20, True) + 60
        return tag(_center(i) - w / 2, y, w, label, color, fg, h, name)

    # Farmacéutico
    px_, py_, pw_, ph_ = PHARM
    image(s, "farmaceutico-xray", px_, py_, pw_, ph_, label="Farmacéutico (radiografía)")
    text(s, px_ + pw_ + 18, py_ + 2, 300, 44, "Farmacéutico", size=20, bold=True, color=WHITE)
    tag(px_ + pw_ + 18, py_ + 52, None, "borrador + revisión obligatoria", soft,
        name="Estado farmacéutico")
    line(s, 960, py_ + ph_ + 2, 960, fy, color=XRAY, lw=2, dash=MSO_LINE.ROUND_DOT,
         tail="triangle", name="Supervisión")

    # Arnés en rojo
    box(s, FRAME_X, fy, FRAME_W, fh, line=ROJO, lw=3, dash=MSO_LINE.DASH, radius=24, name="Arnés")
    box(s, FRAME_X + 26, fy - 22, 120, 44, fill=XRAY_BG, paras=["Arnés"], size=20, bold=True,
        color=WHITE, align="c", anchor="m", name="Rótulo arnés")
    tag(FRAME_X + FRAME_W - 820, fy + fh + 16, 820,
        f"sin tope de dosis · 1 de cada 6 casos >{NB}4{NB}g/día", ROJO, fg=WHITE,
        name="Estado arnés")

    img_y, img_h = fy + 28, 150
    scheme_art(s, "xray", img_y, img_h, XRAY)
    name_y = img_y + img_h + 14
    # Contexto
    text(s, PCOLS[0][0], name_y, PCOLS[0][1], 44, "Contexto", size=20, bold=True, color=WHITE,
         align="c")
    ctag(0, name_y + 52, 360, "guías locales recuperadas", soft, h=84, name="Estado contexto")
    # Modelo
    text(s, PCOLS[1][0], name_y, PCOLS[1][1], 44, "Modelo (GPT-4o)", size=20, bold=True,
         color=WHITE, align="c")
    ctag(1, name_y + 52, None, f"interpreta: 83{NB}%", AMBAR, name="Estado modelo")
    f1 = ctag(1, name_y + 106, None, f"predice concentraciones: 58{NB}%", ROJO, fg=WHITE,
              name="Fallo 1")
    ctag(1, name_y + 160, None, "cuándo extraer: 0 de 30", ROJO, fg=WHITE, name="Fallo 2")
    # Herramienta
    text(s, PCOLS[2][0], name_y, PCOLS[2][1], 140, [
        [("Herramienta", {"bold": True, "color": WHITE})],
        [("(motor de ecuaciones de primer orden)", {"color": "C9D6E2"})],
    ], size=20, align="c")
    ctag(2, name_y + 152, None, "cálculos sin errores", VERDE, name="Estado herramienta")

    # Flechas curvas
    fx_r = _center(1) + (text_width(f"predice concentraciones: 58{NB}%", 20, True) + 60) / 2
    fy_r = name_y + 128
    curve(s, (fx_r + 6, fy_r), (fx_r + 40, fy_r - 70), (fx_r + 50, img_y + img_h - 30),
          (_center(2) - 66, img_y + img_h - 40), color=WHITE, lw=2.5)
    text(s, _center(1) + 62, img_y + img_h - 32, 250, 44, "esto es cálculo", size=20,
         italic=True, color=WHITE, align="l", name="Rótulo cálculo")
    tx_end = FRAME_X + FRAME_W
    curve(s, (tx_end - 24, fy + fh + 16), (tx_end + 40, fy + fh - 10), (tx_end + 60, fy + fh - 120),
          (tx_end + 4, fy + fh - 170), color=WHITE, lw=2.5)
    text(s, tx_end + 14, fy + 120, 140, 100, "esto es un límite", size=20, italic=True,
         color=WHITE, align="l", name="Rótulo límite")

    text(s, MARGIN + 40, 920, 1400, 46,
         "Prepublicación · 30 casos retrospectivos · función renal estable", size=20,
         bold=True, color=AMBAR, align="l", anchor="m", name="Franja")
    source(s, "Hassan et al., medRxiv, 2026 (prepublicación)")


def _freeform(s, pts, fill, name):
    ff = s.shapes.build_freeform(px(pts[0][0]), px(pts[0][1]))
    ff.add_line_segments([(px(x), px(y)) for x, y in pts[1:]], close=True)
    sh = ff.convert_to_shape()
    _strip_style(sh)
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(INK)
    sh.line.width = Pt(2.5)
    _name(sh, name)
    return sh


def slide_label(s):
    title(s, "Leemos la ficha técnica de lo que dispensamos; leamos también la de la IA "
             "que usamos.")
    # Caja de medicamento genérica (formas editables, proyección oblicua)
    fx0, fy0, fw0, fh0 = 96, 420, 480, 500       # cara frontal
    dx, dy = 112, 76
    _freeform(s, [(fx0, fy0), (fx0 + dx, fy0 - dy), (fx0 + fw0 + dx, fy0 - dy), (fx0 + fw0, fy0)],
              "E9EEF2", "Caja: tapa")
    _freeform(s, [(fx0 + fw0, fy0), (fx0 + fw0 + dx, fy0 - dy), (fx0 + fw0 + dx, fy0 + fh0 - dy),
                  (fx0 + fw0, fy0 + fh0)], "D5DDE4", "Caja: lateral")
    lat = text(s, fx0 + fw0 + dx / 2 - 200, fy0 + fh0 / 2 - dy / 2 - 48, 400, 96,
               "Lea la ficha técnica antes de usar", size=20, bold=True, color=INK, align="c",
               anchor="m", name="Caja: texto lateral")
    lat.rotation = 270
    box(s, fx0, fy0, fw0, fh0, fill=WHITE, line=INK, lw=2.5, name="Caja: frontal")
    box(s, fx0 + 1.5, fy0 + 30, fw0 - 3, 18, fill=TURQ, name="Caja: banda")
    box(s, fx0 + 1.5, fy0 + 54, fw0 - 3, 6, fill=LIMA, name="Caja: banda fina")
    text(s, fx0 + 32, fy0 + 84, fw0 - 64, 90, "Dosifica-IA", size=44, bold=True, color=INK,
         anchor="m", name="Caja: nombre")
    text(s, fx0 + 32, fy0 + 180, fw0 - 64, 110, "Contiene: 1 modelo de lenguaje", size=22,
         color=INK, name="Caja: contenido")
    box(s, fx0 + fw0 - 168, fy0 + fh0 - 188, 140, 140, fill=TURQ_T, kind=MSO_SHAPE.OVAL,
        name="Caja: fondo pictograma")
    image(s, "loro-pictograma", fx0 + fw0 - 158, fy0 + fh0 - 180, 120, 124,
          label="Pictograma: loro")
    box(s, fx0 + 32, fy0 + fh0 - 120, 210, 2, fill=RULE, name="Caja: línea")
    box(s, fx0 + 32, fy0 + fh0 - 98, 150, 2, fill=RULE, name="Caja: línea")

    # Ficha técnica desplegada
    lx, lw_ = cols(6, 12)
    ly, lh = 298, 680
    box(s, lx + 10, ly + 10, lw_ - 10, lh - 10, fill="E3E8ED", radius=8, name="Ficha: sombra")
    box(s, lx, ly, lw_ - 10, lh - 10, fill=WHITE, line=RULE, lw=1.5, radius=8, name="Ficha técnica")
    sections = [
        ("icono-dos-columnas", "Composición", "¿Qué hay dentro?",
         " Estadística con modelo, aprendizaje automático o modelo de lenguaje."),
        ("icono-esquema", "Mecanismo de acción", "¿Quién hace el cálculo?",
         " Una herramienta validada o el modelo de lenguaje."),
        ("icono-escalera", "Eficacia clínica y seguridad", "¿Qué ha demostrado, y frente a qué?",
         " Aciertos en una base de datos o resultados en pacientes; frente a un bayesiano con "
         "concentraciones o frente a un rival débil."),
    ]
    sy = [ly + 28, ly + 214, ly + 400]
    for k, ((ic, head, q, ans), y) in enumerate(zip(sections, sy)):
        if k:
            line(s, lx + 24, y - 18, lx + lw_ - 34, y - 18, color="E2E7EC", lw=1.25,
                 dash=MSO_LINE.DASH, name="Pliegue")
        image(s, ic, lx + 34, y + 6, 92, 80, label=f"Icono {head}")
        text(s, lx + 150, y, lw_ - 194, 250, [
            {"runs": [(head, {"bold": True, "color": TURQ_D, "size": 22})], "after": 2},
            [(q, {"bold": True}), (ans, {})],
        ], size=20, color=INK)


def slide_close(s):
    title(s, [[("La ", {}), ("herramienta", {"bold": True, "color": VERDE_D}), (" calcula, el ", {}),
               ("modelo", {"bold": True, "color": MORADO}), (" explica", {})],
              [("y el ", {}), ("farmacéutico", {"bold": True, "color": AZUL_D}), (" decide.", {})]],
          size=36, align="c")
    # Esquema reducido en gris (columnas 3-10); farmacéutico a todo color
    fx, fw = cols(3, 10)
    k = fw / FRAME_W
    pw_, ph_ = 150, 145
    image(s, "farmaceutico-color", 960 - pw_ / 2, 296, pw_, ph_,
          label="Farmacéutico con pluma (a todo color)")
    fy, fh = 470, 380
    line(s, 960, 296 + ph_ + 2, 960, fy, color=AZUL, lw=2.25, dash=MSO_LINE.ROUND_DOT,
         tail="triangle", name="Supervisión")
    box(s, fx, fy, fw, fh, line="C3CBD3", lw=2.25, dash=MSO_LINE.DASH, radius=20,
        name="Arnés (gris)")

    def cx(i):
        return 960 + (_center(i) - 960) * k

    iy, ih = fy + 70, 240
    image(s, "ficha-caso-gray", cx(0) - 80, iy, 160, ih)
    image(s, "loro-gray", cx(1) - 70, iy - 10, 140, ih + 20)
    st = ih / 3
    for i, ic in enumerate(("curva", "campana", "red")):
        image(s, f"icono-{ic}-gray", cx(2) - 50, iy + i * st + 2, 100, st - 8)
    ay = iy + ih / 2
    line(s, cx(0) + 100, ay, cx(1) - 80, ay, color="C3CBD3", lw=2.25, tail="triangle")
    line(s, cx(1) + 80, ay, cx(2) - 80, ay, color="C3CBD3", lw=2.25, head="triangle",
         tail="triangle")

    box(s, MARGIN, 880, 1760, 100, fill=LIGHT, radius=16, name="Relevo", paras=[[
        ("A continuación: ", {"bold": True, "color": TURQ_D}),
        ("Emilio Monte Boquet", {"bold": True}),
        (" · Competencias y habilidades en IA para el farmacéutico hospitalario: de usuario a "
         "supervisor experto", {})]],
        size=20, color=INK, align="l", anchor="m", margins=(28, 4, 28, 4))


# --- Montaje ------------------------------------------------------------------

def build(template, out):
    prs = Presentation(template)
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)

    # Quitar las diapositivas de ejemplo 2-4 (la portada se conserva)
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids)[1:]:
        prs.part.drop_rel(sld_id.get(qn("r:id")))
        sld_ids.remove(sld_id)

    cover = prs.slides[0]
    slide_cover(cover)
    layout = prs.slide_layouts[0]                  # fondo de contenido de la plantilla
    for builder in (slide_hook, slide_not_ai, slide_thirty_years, slide_ladder,
                    slide_pieces, slide_xray, slide_label, slide_close):
        sl = prs.slides.add_slide(layout)
        builder(sl)
    prs.save(out)
    print(f"Escrito {out}")


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
