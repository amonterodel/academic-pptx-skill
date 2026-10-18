"""Inserta una diapositiva nueva en una presentación ya montada, sin tocar las demás.

    python3 insertar_diapositiva.py <entrada.pptx> <salida.pptx> <función> <posición>

<función> es el nombre de una función slide_* de build_deck.py y <posición> el
número de la diapositiva nueva (1 = primera). Ejemplo, el caso de las tres IA
como diapositiva 2:

    python3 insertar_diapositiva.py salida/ponencia-ia-tdm-montero.pptx \
        salida/ponencia-ia-tdm-montero.pptx slide_case 2
"""

import sys

from pptx import Presentation

import build_deck


def insert(src, out, func_name, position):
    prs = Presentation(src)
    builder = getattr(build_deck, func_name)
    slide = prs.slides.add_slide(prs.slide_layouts[0])     # fondo de contenido de la plantilla
    builder(slide)
    ids = prs.slides._sldIdLst
    new_id = list(ids)[-1]
    ids.remove(new_id)
    ids.insert(position - 1, new_id)
    prs.save(out)
    print(f"Escrito {out}: «{func_name}» como diapositiva {position} de {len(prs.slides)}")


if __name__ == "__main__":
    insert(sys.argv[1], sys.argv[2], sys.argv[3], int(sys.argv[4]))
